"""Bezpieczny import XLSX dla POST /upload: skoroszyt -> CSV (UTF-8), dalej ta sama sciezka co CSV od klienta.

Kontrole (kazda konczy sie odmowa z wlasnym kodem powodu, komunikat bez tresci pliku):
    1. kontener ZIP: poprawny archiwum, liczba wpisow, suma rozpakowanych bajtow i wspolczynnik kompresji
       (zip-bomb) liczone z katalogu centralnego PRZED jakimkolwiek parsowaniem XML,
    2. makra i obiekty osadzone: `vbaProject.bin`, `macrosheets/`, `embeddings/`, typ zawartosci `macroEnabled`,
    3. DTD/encje w CALEJ czesci XML (skan strumieniowy, odmowa przy pierwszym `<!DOCTYPE`/`<!ENTITY`; XXE,
       billion laughs) oraz kodowania UTF-16/32 omijajace skan; dodatkowo `openpyxl` parsuje przez `defusedxml`,
    4. arkusze: dokladnie jeden albo wskazany polem `sheet`; twardy limit liczby arkuszy,
    5. odczyt `openpyxl` w trybie `read_only=True, data_only=True` (formuly nie sa liczone ani przepisywane;
       zbuforowana wartosc albo puste pole), z `reset_dimensions()` — deklarowany wymiar arkusza jest
       niezaufany (`A1:XFD1048576` rozdmuchalby kazdy wiersz do 16384 komorek),
    6. limity wierszy i kolumn,
    7. limit `sharedStrings` (liczba wpisow i rozmiar) oraz limit czasu calej konwersji (`XLSX_PARSE_TIMEOUT_SEC`).

Odrzucamy rowniez skoroszyty z makrami zapisane pod nazwa `.xlsx`; `.xlsm`/`.xls`/`.xlsb` odpadaja juz na
rozszerzeniu (`upload_guard.validate_original_filename`). Zewnetrzne powiazania (`externalLinks`) nie sa
sledzone: tryb `data_only` czyta tylko zbuforowane wartosci.
"""
from __future__ import annotations

import csv
import io
import logging
import re
import threading
import zipfile
from collections.abc import Callable
from concurrent.futures import ThreadPoolExecutor, wait
from datetime import date, datetime, time, timedelta
from pathlib import Path
from typing import Any, TypeVar

import openpyxl

from app.core.upload_guard import ReasonCode, UploadRejected

log = logging.getLogger(__name__)

T = TypeVar("T")

# --------------------------------------------------------------------------- limity
MAX_XLSX_ROWS = 200_000                          # wierszy danych (bez naglowka)
MAX_XLSX_COLUMNS = 256                           # kolumn w naglowku/wierszu (po obcieciu pustych koncowek)
MAX_XLSX_SHEETS = 20                             # arkuszy w skoroszycie, niezaleznie od pola `sheet`
MAX_XLSX_ENTRIES = 5_000                         # wpisow w archiwum ZIP
MAX_XLSX_UNCOMPRESSED_BYTES = 48 * 1024 * 1024   # suma rozpakowanych bajtow wszystkich wpisow
MAX_XLSX_RATIO = 100                             # rozpakowane / skompresowane, liczone od RATIO_CHECK_FROM_BYTES
RATIO_CHECK_FROM_BYTES = 10 * 1024 * 1024
MAX_XLSX_SCANNED_ROWS = 1_048_577                # twardy limit Excela (+ naglowek); chroni czas odczytu
MAX_SHEET_NAME_LEN = 64
MAX_SHARED_STRINGS = 200_000                     # wpisow <si> w xl/sharedStrings.xml
MAX_SHARED_STRINGS_BYTES = 24 * 1024 * 1024      # rozmiar rozpakowanego xl/sharedStrings.xml
XLSX_PARSE_TIMEOUT_SEC = 30.0                    # caly czas konwersji (kontener + odczyt arkusza)
XLSX_CANCEL_GRACE_SEC = 2.0                      # po timeoucie: ile czekamy, az watek zwolni uchwyt pliku
_SCAN_CHUNK_BYTES = 256 * 1024
_CANCEL_CHECK_EVERY_ROWS = 500
_MAX_CONTENT_TYPES_BYTES = 1024 * 1024
_SHARED_STRINGS_PART = "xl/sharedstrings.xml"
_SI_TAG = re.compile(rb"<(?:\w+:)?si[\s/>]")
_XML_BOMS = (b"\xff\xfe", b"\xfe\xff", b"\x00\x00\xfe\xff", b"\xff\xfe\x00\x00")
_XML_FORBIDDEN_MARKERS = (b"<!DOCTYPE", b"<!ENTITY")
_OVERLAP = 64                                    # bajty na styku porcji: znacznik/tag nie ucieknie skanowi

#: Przedrostki nazw wpisow ZIP swiadczace o makrach lub obiektach osadzonych (porownanie po lower()).
_FORBIDDEN_PREFIXES = ("xl/macrosheets/", "xl/dialogsheets/", "xl/embeddings/", "xl/activex/")
_FORBIDDEN_SUFFIXES = ("vbaproject.bin", "vbaprojectsignature.bin", "vbadata.xml")

def _reject(reason: str, status: int, message: str) -> UploadRejected:
    return UploadRejected(reason, status, message)


def _invalid() -> UploadRejected:
    return _reject(ReasonCode.XLSX_INVALID, 415, "plik nie jest poprawnym skoroszytem XLSX")


def _too_large() -> UploadRejected:
    return _reject(ReasonCode.XLSX_UNPACKED_TOO_LARGE, 413, "skoroszyt po rozpakowaniu jest zbyt duzy")


class _Cancelled(Exception):
    """Wewnetrzny sygnal: wywolujacy przestal czekac (timeout), watek konwersji ma sie zakonczyc."""


def _check_cancel(cancel: threading.Event | None) -> None:
    if cancel is not None and cancel.is_set():
        raise _Cancelled


def _scan_xml_part(part: Any, *, count_strings: bool, cancel: threading.Event | None) -> None:
    """Skan strumieniowy CALEJ czesci XML: odmowa przy pierwszym `<!DOCTYPE`/`<!ENTITY` albo BOM UTF-16/32.

    Porcje sa czytane z nakladka `_OVERLAP`, wiec znacznik na granicy porcji nie umyka skanowi. Dla
    `sharedStrings` liczy tez wpisy `<si>` i rozmiar (limity `MAX_SHARED_STRINGS*`).
    """
    carry = b""
    first = True
    size = 0
    strings = 0
    while True:
        _check_cancel(cancel)
        chunk = part.read(_SCAN_CHUNK_BYTES)
        if not chunk:
            return
        if first:
            first = False
            if chunk.startswith(_XML_BOMS):
                raise _invalid()  # UTF-16/32 omijaloby skan DTD
        window = carry + chunk
        if any(marker in window for marker in _XML_FORBIDDEN_MARKERS):
            raise _invalid()
        if count_strings:
            size += len(chunk)
            strings += sum(1 for m in _SI_TAG.finditer(window) if m.end() > len(carry))
            if size > MAX_SHARED_STRINGS_BYTES or strings > MAX_SHARED_STRINGS:
                raise _too_large()
        carry = window[-_OVERLAP:]


# --------------------------------------------------------------------------- kontener
def validate_container(path: Path, cancel: threading.Event | None = None) -> None:
    """Sprawdza archiwum ZIP bez parsowania zawartosci arkuszy. Rzuca `UploadRejected`.

    `cancel` (opcjonalne) przerywa skan, gdy wywolujacy przestal czekac (patrz `xlsx_to_csv_bytes`).
    """
    try:
        with path.open("rb") as fh, zipfile.ZipFile(fh) as zf:
            infos = zf.infolist()
            if not infos or len(infos) > MAX_XLSX_ENTRIES:
                raise _invalid()

            total = sum(i.file_size for i in infos)
            packed = max(1, sum(i.compress_size for i in infos))
            if total > MAX_XLSX_UNCOMPRESSED_BYTES or (
                    total > RATIO_CHECK_FROM_BYTES and total / packed > MAX_XLSX_RATIO):
                raise _too_large()

            names = {i.filename.lower(): i for i in infos}
            if "[content_types].xml" not in names or "xl/workbook.xml" not in names:
                raise _invalid()
            if any(i.flag_bits & 0x1 for i in infos):  # wpis zaszyfrowany
                raise _invalid()

            for name in names:
                if name.startswith(_FORBIDDEN_PREFIXES) or name.endswith(_FORBIDDEN_SUFFIXES):
                    raise _reject(ReasonCode.XLSX_MACROS, 415,
                                  "skoroszyty z makrami i obiektami osadzonymi nie sa przyjmowane")

            ct = names["[content_types].xml"]
            if ct.file_size > _MAX_CONTENT_TYPES_BYTES:
                raise _invalid()
            if b"macroenabled" in zf.read(ct).lower():
                raise _reject(ReasonCode.XLSX_MACROS, 415,
                              "skoroszyty z makrami i obiektami osadzonymi nie sa przyjmowane")

            shared = names.get(_SHARED_STRINGS_PART)
            if shared is not None and shared.file_size > MAX_SHARED_STRINGS_BYTES:
                raise _too_large()

            for info in infos:  # kazdy wpis, rowniez zdublowane nazwy (nie tylko ostatni z `names`)
                name = info.filename.lower()
                if not name.endswith((".xml", ".rels")):
                    continue
                with zf.open(info) as part:
                    _scan_xml_part(part, count_strings=name == _SHARED_STRINGS_PART, cancel=cancel)
    except (UploadRejected, _Cancelled):
        raise
    except (zipfile.BadZipFile, NotImplementedError, RuntimeError, OSError, EOFError, ValueError) as exc:
        log.info("xlsx container rejected: %s", type(exc).__name__)
        raise _invalid() from None


# --------------------------------------------------------------------------- komorki
def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        return format(value, ".15g")  # 0.1+0.2 -> "0.3", 899.0 -> "899"
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d") if value.time() == time.min else value.strftime("%Y-%m-%d %H:%M:%S")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, (time, timedelta)):
        return value.isoformat() if isinstance(value, time) else str(value)
    return str(value)


def _trim(values: tuple[Any, ...]) -> list[str]:
    cells = [_cell_text(v) for v in values]
    while cells and not cells[-1].strip():
        cells.pop()
    return cells


# --------------------------------------------------------------------------- konwersja
def xlsx_to_csv_bytes(path: Path, sheet: str | None = None) -> tuple[bytes, int]:
    """Waliduje skoroszyt i zwraca (CSV w UTF-8, liczba wierszy danych). Rzuca `UploadRejected`.

    Cala praca (kontener + odczyt arkusza) idzie w osobnym watku, na ktory czekamy co najwyzej
    `XLSX_PARSE_TIMEOUT_SEC`; po przekroczeniu: 422 `xlsx_timeout`. Watku nie da sie zabic, wiec dostaje
    flage `cancel` sprawdzana w petlach skanu i odczytu wierszy i konczy sie sam; czekamy na to najwyzej
    `XLSX_CANCEL_GRACE_SEC`, zeby plik tymczasowy byl juz zamkniety, gdy handler go usuwa.
    """
    return run_bounded(_convert, path, sheet)


def run_bounded(func: Callable[..., T], *args: Any) -> T:
    """Wykonuje `func(*args, cancel)` w watku, czekajac najwyzej `XLSX_PARSE_TIMEOUT_SEC` (inaczej `xlsx_timeout`).

    Wspolny limit czasu dla konwersji i dla skanu formul (`xlsx_formulas`). `func` dostaje na koncu
    `threading.Event`, ktory zostaje ustawiony po przekroczeniu czasu.
    """
    cancel = threading.Event()
    pool = ThreadPoolExecutor(max_workers=1, thread_name_prefix="xlsx-parse")
    try:
        future = pool.submit(func, *args, cancel)
        done, _ = wait([future], timeout=XLSX_PARSE_TIMEOUT_SEC)
        if not done:
            cancel.set()
            # Krotka laska: watek sprawdza `cancel` co porcje/500 wierszy i zamyka plik sam. Bez tego wywolujacy
            # usuwalby plik tymczasowy, ktory jest jeszcze otwarty — na Windows to blad wspoldzielenia (WinError 32).
            released, _ = wait([future], timeout=XLSX_CANCEL_GRACE_SEC)
            log.warning("xlsx parse timed out after %.1fs (worker released: %s)", XLSX_PARSE_TIMEOUT_SEC,
                        bool(released))
            raise _reject(ReasonCode.XLSX_TIMEOUT, 422, "odczyt skoroszytu trwal zbyt dlugo")
        return future.result()
    finally:
        pool.shutdown(wait=False, cancel_futures=True)


def _convert(path: Path, sheet: str | None, cancel: threading.Event) -> tuple[bytes, int]:
    validate_container(path, cancel)

    wb = None
    try:
        with path.open("rb") as fh:  # uchwyt zamiast sciezki: openpyxl sprawdza rozszerzenie nazwy, a plik to *.tmp
            wb = openpyxl.load_workbook(fh, read_only=True, data_only=True, keep_links=False)
            names = list(wb.sheetnames)
            if not names:
                raise _invalid()
            if len(names) > MAX_XLSX_SHEETS:
                raise _reject(ReasonCode.XLSX_TOO_MANY_SHEETS, 422, "skoroszyt ma zbyt wiele arkuszy")
            if sheet is not None:
                if sheet not in names:
                    raise _reject(ReasonCode.XLSX_SHEET_NOT_FOUND, 422, "nie znaleziono wskazanego arkusza")
                chosen = sheet
            elif len(names) > 1:
                raise _reject(ReasonCode.XLSX_MULTIPLE_SHEETS, 422,
                              "skoroszyt ma wiele arkuszy; wskaz arkusz w polu sheet")
            else:
                chosen = names[0]

            ws = wb[chosen]
            ws.reset_dimensions()  # niezaufany wymiar z pliku (patrz docstring modulu)
            return _sheet_to_csv(ws, cancel)
    except (UploadRejected, _Cancelled):
        raise
    except Exception as exc:  # noqa: BLE001  openpyxl/XML rzucaja wiele typow; do klienta idzie tylko komunikat ogolny
        log.info("xlsx parse rejected: %s", type(exc).__name__, exc_info=log.isEnabledFor(logging.DEBUG))
        raise _invalid() from None
    finally:
        if wb is not None:
            wb.close()


def _sheet_to_csv(ws: Any, cancel: threading.Event | None = None) -> tuple[bytes, int]:
    out = io.StringIO()
    writer = csv.writer(out, lineterminator="\n")
    header: list[str] | None = None
    width = 0
    data_rows = 0

    for scanned, values in enumerate(ws.iter_rows(values_only=True), start=1):
        if scanned % _CANCEL_CHECK_EVERY_ROWS == 0:
            _check_cancel(cancel)
        if scanned > MAX_XLSX_SCANNED_ROWS:
            raise _reject(ReasonCode.TOO_MANY_RECORDS, 413, "arkusz zawiera zbyt wiele wierszy")
        cells = _trim(values)
        if not cells:
            continue  # pusty wiersz (takze poczatkowe puste wiersze przed naglowkiem)
        if len(cells) > MAX_XLSX_COLUMNS:
            raise _reject(ReasonCode.XLSX_TOO_WIDE, 413, "arkusz ma zbyt wiele kolumn")
        if header is None:
            header, width = cells, len(cells)
            writer.writerow(header)
            continue
        data_rows += 1
        if data_rows > MAX_XLSX_ROWS:
            raise _reject(ReasonCode.TOO_MANY_RECORDS, 413, f"arkusz zawiera wiecej niz {MAX_XLSX_ROWS} wierszy")
        writer.writerow((cells + [""] * width)[:width])  # kolumny bez naglowka sa pomijane

    if header is None:
        raise _reject(ReasonCode.HEADER_MISSING, 422, "arkusz nie zawiera naglowka")
    return out.getvalue().encode("utf-8"), data_rows
