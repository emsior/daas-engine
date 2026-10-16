"""Wykrywanie formul XLSX bez zbuforowanych wartosci (C14).

`xlsx_import` czyta skoroszyt w trybie `data_only=True`: komorka z formula, ktorej Excel/LibreOffice nie zapisal
wyniku (plik wygenerowany biblioteka, eksport "bez przeliczenia"), wychodzi jako PUSTA — a wiec np. przychod znika
z wiersza bez sladu. Ten modul liczy takie komorki, zeby wywolujacy mogl ostrzec uzytkownika.

Funkcja zaklada plik juz zwalidowany przez `xlsx_import.validate_container` (limity, brak makr, brak DTD).
"""
from __future__ import annotations

import logging
import threading
from pathlib import Path
from typing import Any

import openpyxl

from app.core import xlsx_import
from app.core.upload_guard import UploadRejected

log = logging.getLogger(__name__)

FORMULA_WARNING = ("{n} komórek z formułami bez zapisanych wartości — zapisz plik ponownie w Excelu/LibreOffice "
                   "(formuły bez wartości zostały odczytane jako puste)")


def _sheet(wb: Any, sheet: str | None) -> Any:
    name = sheet if sheet is not None else wb.sheetnames[0]
    ws = wb[name]
    ws.reset_dimensions()  # wymiar z pliku jest niezaufany (jak w xlsx_import)
    return ws


def count_formulas_without_value(path: Path, sheet: str | None = None,
                                 cancel: threading.Event | None = None) -> int:
    """Liczba komorek z formula, dla ktorych brak zbuforowanej wartosci (pierwszy arkusz albo `sheet`).

    Blad odczytu zwraca 0 i trafia do logu — konwersja pliku (`xlsx_to_csv_bytes`) odrzuci taki plik osobno,
    wiec ostrzezenie nie moze go zasloniac. `cancel` (ustawiany po przekroczeniu limitu czasu) przerywa skan: 0.
    """
    try:
        with path.open("rb") as fh_f, path.open("rb") as fh_v:
            wb_f = openpyxl.load_workbook(fh_f, read_only=True, data_only=False, keep_links=False)
            wb_v = openpyxl.load_workbook(fh_v, read_only=True, data_only=True, keep_links=False)
            try:
                rows_f = _sheet(wb_f, sheet).iter_rows(values_only=True)
                rows_v = _sheet(wb_v, sheet).iter_rows(values_only=True)
                missing = 0
                for scanned, (row_f, row_v) in enumerate(zip(rows_f, rows_v, strict=False), start=1):
                    if scanned > xlsx_import.MAX_XLSX_SCANNED_ROWS:
                        break
                    if cancel is not None and scanned % 500 == 0 and cancel.is_set():
                        return 0
                    missing += sum(1 for f, v in zip(row_f, row_v, strict=False)
                                   if isinstance(f, str) and f.startswith("=") and v is None)
                return missing
            finally:
                wb_f.close()
                wb_v.close()
    except Exception as exc:  # noqa: BLE001  openpyxl rzuca wiele typow; szczegol tylko w logu
        log.warning("formula scan failed: %s", type(exc).__name__)
        return 0


def formula_warning(path: Path, sheet: str | None = None) -> str | None:
    """Gotowe ostrzezenie dla uzytkownika albo `None`, gdy wszystkie formuly maja wartosci.

    Skan jest ograniczony tym samym limitem czasu co konwersja (`xlsx_import.run_bounded`). Po przekroczeniu
    czasu ostrzezenia nie ma (upload i tak przeszedl walidacje kontenera): to podpowiedz, nie warunek przyjecia.
    """
    try:
        n = xlsx_import.run_bounded(count_formulas_without_value, path, sheet)
    except UploadRejected:
        log.warning("formula scan skipped: time limit")
        return None
    return FORMULA_WARNING.format(n=n) if n else None
