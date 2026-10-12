"""POST /upload przyjmuje XLSX z limitami (wariant B): bez makr, jeden arkusz (albo `sheet`), limity wierszy,
rozpakowanych bajtow i wpisow. Plik jest konwertowany do CSV i idzie ta sama sciezka co CSV od klienta."""
from __future__ import annotations

import hashlib
import io
import json
import zipfile
from datetime import datetime
from pathlib import Path

import openpyxl
import pytest
from fastapi.testclient import TestClient

import app.core.xlsx_import as xi
from app.core.config import Settings
from app.core.upload_guard import AUDIT_ALLOWED_KEYS, MAX_UPLOAD_BYTES, ReasonCode
from app.sources.csv_mapper import read_client_csv
from tests.conftest import PROJECT_ROOT

PL_FIXTURE = PROJECT_ROOT / "data" / "fixtures" / "client_export_pl.csv"
XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


# ---------------------------------------------------------------- budowanie plikow
def workbook_bytes(wb: openpyxl.Workbook) -> bytes:
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def fixture_workbook() -> openpyxl.Workbook:
    """Eksport z fixtures/client_export_pl.csv jako tekstowe komorki (tak zapisuje je wiekszosc sklepow)."""
    df = read_client_csv(str(PL_FIXTURE))
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Zamowienia"
    ws.append(list(df.columns))
    for row in df.itertuples(index=False):
        ws.append([str(v) for v in row])
    return wb


def rewrite_zip(raw: bytes, edit) -> bytes:
    """Kopiuje archiwum, pozwalajac `edit(name, data) -> (name, data) | None` zmienic/usunac wpis."""
    out = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(raw)) as zin, zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zout:
        for info in zin.infolist():
            res = edit(info.filename, zin.read(info.filename))
            if res is not None:
                zout.writestr(res[0], res[1])
    return out.getvalue()


def add_entry(raw: bytes, name: str, data: bytes) -> bytes:
    out = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(raw)) as zin, zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zout:
        for info in zin.infolist():
            zout.writestr(info, zin.read(info.filename))
        zout.writestr(name, data)
    return out.getvalue()


def post(client: TestClient, raw: bytes, name: str = "eksport.xlsx", **data):
    return client.post("/upload", files={"file": (name, raw, XLSX_MIME)}, data=data)


def uploads(settings: Settings) -> Path:
    return settings.reports_path.parent / "uploads"


def assert_clean(settings: Settings) -> None:
    up = uploads(settings)
    if up.exists():
        assert list(up.glob("*.csv")) == [] and list(up.glob("*.tmp")) == [] and list(up.glob("*.xlsx")) == []


def audit(settings: Settings) -> list[dict]:
    p = settings.audit_log_path
    recs = [json.loads(x) for x in p.read_text(encoding="utf-8").splitlines() if x.strip()] if p.exists() else []
    return [r for r in recs if r.get("action") == "upload"]


# ---------------------------------------------------------------- poprawny plik
def test_valid_xlsx_goes_through_the_csv_path(client: TestClient, settings: Settings):
    raw = workbook_bytes(fixture_workbook())
    r = post(client, raw, client_name="Sklep XLSX")
    assert r.status_code == 200, r.text
    up = r.json()
    assert up["ready"] is True and up["rows"] == 15 and up["confidence"] >= 0.8
    assert up["filename"].endswith(".csv")  # na dysku zostaje CSV, nie skoroszyt
    assert [p.suffix for p in uploads(settings).iterdir()] == [".csv"]

    run = client.post("/run", json=up["run_payload"])
    assert run.status_code == 200 and run.json()["run_status"] == "SUCCESS"
    assert run.json()["data_status"] == "LIVE_DATA" and run.json()["client_name"] == "Sklep XLSX"


def test_xlsx_and_csv_of_the_same_data_give_the_same_metrics(client: TestClient):
    csv_up = client.post("/upload", files={"file": ("a.csv", PL_FIXTURE.read_bytes(), "text/csv")}).json()
    xlsx_up = post(client, workbook_bytes(fixture_workbook())).json()
    m_csv = client.post("/run", json=csv_up["run_payload"]).json()["metrics"]
    m_xlsx = client.post("/run", json=xlsx_up["run_payload"]).json()["metrics"]
    assert m_csv["revenue_total"] == m_xlsx["revenue_total"] and m_csv["orders_total"] == m_xlsx["orders_total"]
    assert m_csv["margin_pct"] == m_xlsx["margin_pct"]


def test_typed_cells_numbers_and_dates_are_converted(client: TestClient):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Nr zamówienia", "Data sprzedaży", "Nazwa produktu", "Kategoria", "Ilość", "Wartość brutto",
               "Koszt zakupu", "Status zamówienia"])
    ws.append(["Z-1", datetime(2026, 8, 3), "Fotel", "Meble", 2, 1299.5, 870, "Zrealizowane"])
    ws.append(["Z-2", datetime(2026, 8, 4, 13, 30), "Lampka", "Akcesoria", 1, 0.1 + 0.2, True, "Zrealizowane"])
    r = post(client, workbook_bytes(wb))
    assert r.status_code == 200, r.text
    run = client.post("/run", json=r.json()["run_payload"]).json()
    assert run["run_status"] == "SUCCESS" and run["records_processed"] == 2
    assert run["metrics"]["revenue_total"] == pytest.approx(1299.8)


def test_formulas_are_not_evaluated_or_stored_as_text(client: TestClient, settings: Settings):
    wb = fixture_workbook()
    wb.active["F2"] = "=SUM(1,2)"  # bez zbuforowanej wartosci: data_only=True daje pusta komorke
    r = post(client, workbook_bytes(wb))
    assert r.status_code == 200, r.text
    stored = next(uploads(settings).glob("*.csv")).read_text(encoding="utf-8")
    assert "=SUM" not in stored


def test_crafted_sheet_dimension_does_not_inflate_rows(client: TestClient):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Data", "Kwota", "Produkt"])
    ws.append(["2026-08-01", "100,00", "Kabel"])
    raw = rewrite_zip(workbook_bytes(wb), lambda n, d: (n, d.replace(
        b'<dimension ref="A1:C2" />', b'<dimension ref="A1:XFD1048576" />')))
    r = post(client, raw)
    assert r.status_code == 200, r.text
    assert r.json()["columns"] == ["Data", "Kwota", "Produkt"] and r.json()["rows"] == 1


# ---------------------------------------------------------------- arkusze
def two_sheet_workbook() -> openpyxl.Workbook:
    wb = fixture_workbook()
    junk = wb.create_sheet("Notatki")
    junk.append(["nic", "tu", "nie ma"])
    return wb


def test_multiple_sheets_without_selection_are_rejected(client: TestClient, settings: Settings):
    r = post(client, workbook_bytes(two_sheet_workbook()))
    assert r.status_code == 422
    assert "arkusz" in r.json()["detail"] and "Notatki" not in r.json()["detail"]
    assert audit(settings)[-1]["reason_code"] == "xlsx_multiple_sheets"
    assert_clean(settings)


def test_sheet_field_selects_one_of_many(client: TestClient):
    r = post(client, workbook_bytes(two_sheet_workbook()), sheet="Zamowienia")
    assert r.status_code == 200, r.text
    assert r.json()["rows"] == 15


def test_unknown_sheet_is_rejected(client: TestClient, settings: Settings):
    r = post(client, workbook_bytes(two_sheet_workbook()), sheet="Nie ma takiego")
    assert r.status_code == 422
    assert audit(settings)[-1]["reason_code"] == "xlsx_sheet_not_found"
    assert_clean(settings)


def test_too_many_sheets_are_rejected_even_with_selection(client: TestClient, settings: Settings,
                                                          monkeypatch: pytest.MonkeyPatch):
    monkeypatch.setattr(xi, "MAX_XLSX_SHEETS", 3)
    wb = fixture_workbook()
    for i in range(3):
        wb.create_sheet(f"S{i}")
    r = post(client, workbook_bytes(wb), sheet="Zamowienia")
    assert r.status_code == 422
    assert audit(settings)[-1]["reason_code"] == "xlsx_too_many_sheets"
    assert_clean(settings)


def test_overlong_sheet_name_is_rejected_by_validation(client: TestClient):
    assert post(client, workbook_bytes(fixture_workbook()), sheet="x" * 65).status_code == 422


# ---------------------------------------------------------------- makra i rozszerzenia
@pytest.mark.parametrize("name", ["eksport.xlsm", "eksport.xls", "eksport.xltx", "eksport.xlsb", "eksport.xlsx.exe"])
def test_other_spreadsheet_extensions_are_rejected(client: TestClient, settings: Settings, name: str):
    r = post(client, workbook_bytes(fixture_workbook()), name=name)
    assert r.status_code == 415
    assert_clean(settings)


def test_uppercase_xlsx_extension_is_accepted(client: TestClient):
    assert post(client, workbook_bytes(fixture_workbook()), name="EKSPORT.XLSX").status_code == 200


def test_vba_project_inside_an_xlsx_is_rejected(client: TestClient, settings: Settings):
    raw = add_entry(workbook_bytes(fixture_workbook()), "xl/vbaProject.bin", b"\xd0\xcf\x11\xe0 fake vba")
    r = post(client, raw)
    assert r.status_code == 415
    assert audit(settings)[-1]["reason_code"] == "xlsx_macros_not_allowed"
    assert_clean(settings)


def test_macro_enabled_content_type_is_rejected(client: TestClient, settings: Settings):
    raw = rewrite_zip(workbook_bytes(fixture_workbook()), lambda n, d: (n, d.replace(
        b"application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml",
        b"application/vnd.ms-excel.sheet.macroEnabled.main+xml")) if n == "[Content_Types].xml" else (n, d))
    r = post(client, raw)
    assert r.status_code == 415
    assert audit(settings)[-1]["reason_code"] == "xlsx_macros_not_allowed"
    assert_clean(settings)


@pytest.mark.parametrize("entry", ["xl/macrosheets/sheet1.xml", "xl/embeddings/oleObject1.bin"])
def test_macro_sheets_and_embedded_objects_are_rejected(client: TestClient, settings: Settings, entry: str):
    r = post(client, add_entry(workbook_bytes(fixture_workbook()), entry, b"<x/>"))
    assert r.status_code == 415
    assert_clean(settings)


# ---------------------------------------------------------------- kontener: zip-bomb, DTD, uszkodzenia
def zeros_zip(megabytes: int, extra_valid_parts: bool = True) -> bytes:
    """Poprawnie wygladajacy skoroszyt z jednym ogromnym (skompresowanym do kilku KB) wpisem."""
    base = workbook_bytes(fixture_workbook())
    out = io.BytesIO()
    chunk = b"\0" * (1024 * 1024)
    with zipfile.ZipFile(io.BytesIO(base)) as zin, zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zout:
        if extra_valid_parts:
            for info in zin.infolist():
                zout.writestr(info, zin.read(info.filename))
        with zout.open("xl/media/bomb.bin", "w", force_zip64=True) as fh:
            for _ in range(megabytes):
                fh.write(chunk)
    return out.getvalue()


def test_zip_bomb_is_rejected_before_parsing(client: TestClient, settings: Settings):
    raw = zeros_zip(megabytes=200)  # 200 MiB po rozpakowaniu, ~200 KB w pliku
    assert len(raw) < MAX_UPLOAD_BYTES
    r = post(client, raw)
    assert r.status_code == 413
    assert audit(settings)[-1]["reason_code"] == "xlsx_unpacked_too_large"
    assert_clean(settings)


def test_high_compression_ratio_is_rejected_below_the_absolute_limit(client: TestClient, settings: Settings,
                                                                    monkeypatch: pytest.MonkeyPatch):
    monkeypatch.setattr(xi, "RATIO_CHECK_FROM_BYTES", 1024 * 1024)
    raw = zeros_zip(megabytes=30)  # 30 MiB < limit bezwzgledny, ale ~1000:1
    assert 30 * 1024 * 1024 < xi.MAX_XLSX_UNCOMPRESSED_BYTES
    r = post(client, raw)
    assert r.status_code == 413
    assert audit(settings)[-1]["reason_code"] == "xlsx_unpacked_too_large"


def test_unpacked_limit_is_enforced_per_configured_value(client: TestClient, monkeypatch: pytest.MonkeyPatch):
    monkeypatch.setattr(xi, "MAX_XLSX_UNCOMPRESSED_BYTES", 5 * 1024)
    assert post(client, workbook_bytes(fixture_workbook())).status_code == 413


def test_too_many_zip_entries_are_rejected(client: TestClient, settings: Settings, monkeypatch: pytest.MonkeyPatch):
    monkeypatch.setattr(xi, "MAX_XLSX_ENTRIES", 5)
    r = post(client, workbook_bytes(fixture_workbook()))
    assert r.status_code == 415
    assert audit(settings)[-1]["reason_code"] == "xlsx_invalid"
    assert_clean(settings)


@pytest.mark.parametrize("part", ["docProps/core.xml", "xl/worksheets/sheet1.xml", "xl/_rels/workbook.xml.rels"])
def test_doctype_in_xml_parts_is_rejected(client: TestClient, settings: Settings, part: str):
    evil = (b'<?xml version="1.0"?><!DOCTYPE x [<!ENTITY a "aaaa"><!ENTITY b "&a;&a;&a;&a;">]>'
            b'<root xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><t>&b;</t></root>')
    raw = rewrite_zip(workbook_bytes(fixture_workbook()), lambda n, d: (n, evil) if n == part else (n, d))
    r = post(client, raw)
    assert r.status_code == 415
    assert audit(settings)[-1]["reason_code"] == "xlsx_invalid"
    assert_clean(settings)


def test_utf16_xml_part_is_rejected(client: TestClient, settings: Settings):
    raw = rewrite_zip(workbook_bytes(fixture_workbook()),
                      lambda n, d: (n, "<?xml version='1.0'?><r/>".encode("utf-16")) if n == "docProps/app.xml" else (n, d))
    assert post(client, raw).status_code == 415
    assert_clean(settings)


@pytest.mark.parametrize("payload", [
    b"to jest zwykly tekst, nie archiwum",
    b"PK\x03\x04 uciety naglowek",
    b"",
])
def test_not_a_zip_is_rejected(client: TestClient, settings: Settings, payload: bytes):
    r = post(client, payload or b"x")
    assert r.status_code == 415
    assert_clean(settings)


def test_csv_text_with_xlsx_extension_is_rejected(client: TestClient, settings: Settings):
    r = post(client, PL_FIXTURE.read_bytes())
    assert r.status_code == 415
    assert audit(settings)[-1]["reason_code"] == "xlsx_invalid"
    assert_clean(settings)


def test_truncated_xlsx_is_rejected(client: TestClient, settings: Settings):
    raw = workbook_bytes(fixture_workbook())
    assert post(client, raw[: len(raw) // 2]).status_code == 415
    assert_clean(settings)


def test_zip_without_workbook_part_is_rejected(client: TestClient, settings: Settings):
    raw = rewrite_zip(workbook_bytes(fixture_workbook()), lambda n, d: None if n == "xl/workbook.xml" else (n, d))
    assert post(client, raw).status_code == 415
    assert_clean(settings)


def test_corrupt_sheet_xml_gives_a_generic_error_without_content(client: TestClient, settings: Settings):
    raw = rewrite_zip(workbook_bytes(fixture_workbook()),
                      lambda n, d: (n, b"<worksheet><sheetData><row><c>SEKRET-TRESC") if n == "xl/worksheets/sheet1.xml" else (n, d))
    r = post(client, raw)
    assert r.status_code in (415, 422)
    assert "SEKRET" not in r.text
    assert_clean(settings)


def test_upload_over_the_byte_limit_is_rejected_while_streaming(client: TestClient, settings: Settings):
    r = post(client, b"PK\x03\x04" + b"\0" * MAX_UPLOAD_BYTES)
    assert r.status_code == 413
    assert audit(settings)[-1]["reason_code"] == ReasonCode.FILE_TOO_LARGE
    assert_clean(settings)


# ---------------------------------------------------------------- limity wierszy i szerokosci
def numbered_workbook(rows: int, cols: int = 3) -> openpyxl.Workbook:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Data", "Kwota", "Produkt"] + [f"x{i}" for i in range(cols - 3)])
    for i in range(rows):
        ws.append(["2026-08-01", 10 + i, f"P{i}"] + [1] * (cols - 3))
    return wb


def test_row_limit_constant_matches_the_contract():
    assert xi.MAX_XLSX_ROWS == 200_000


def test_rows_above_the_limit_are_rejected_and_at_the_limit_accepted(client: TestClient, settings: Settings,
                                                                   monkeypatch: pytest.MonkeyPatch):
    monkeypatch.setattr(xi, "MAX_XLSX_ROWS", 50)
    ok = post(client, workbook_bytes(numbered_workbook(50)))
    assert ok.status_code == 200 and ok.json()["rows"] == 50
    for p in uploads(settings).glob("*.csv"):
        p.unlink()
    r = post(client, workbook_bytes(numbered_workbook(51)))
    assert r.status_code == 413
    assert audit(settings)[-1]["reason_code"] == ReasonCode.TOO_MANY_RECORDS
    assert_clean(settings)


def test_xlsx_rows_may_exceed_the_csv_record_limit(monkeypatch: pytest.MonkeyPatch):
    from app.core.upload_guard import MAX_RECORDS, validate_csv_bytes

    body = "Data,Kwota\n" + "2026-08-01,1\n" * (MAX_RECORDS + 10)
    with pytest.raises(Exception, match="rekord"):
        validate_csv_bytes(body.encode())
    assert validate_csv_bytes(body.encode(), max_records=MAX_RECORDS + 10).record_count == MAX_RECORDS + 10


def test_too_wide_rows_are_rejected(client: TestClient, settings: Settings, monkeypatch: pytest.MonkeyPatch):
    monkeypatch.setattr(xi, "MAX_XLSX_COLUMNS", 6)
    r = post(client, workbook_bytes(numbered_workbook(3, cols=8)))
    assert r.status_code == 413
    assert audit(settings)[-1]["reason_code"] == "xlsx_too_wide"
    assert_clean(settings)


# ---------------------------------------------------------------- audyt i polityka
def test_xlsx_upload_is_audited_with_the_hash_of_the_received_file(client: TestClient, settings: Settings):
    raw = workbook_bytes(fixture_workbook())
    assert post(client, raw).status_code == 200
    ev = audit(settings)[-1]
    assert ev["decision"] == "allow" and ev["size_bytes"] == len(raw)
    assert ev["content_sha256"] == hashlib.sha256(raw).hexdigest()
    assert set(ev) <= AUDIT_ALLOWED_KEYS


def test_xlsx_is_accepted_with_policy_enforced(client: TestClient, runner):
    runner.settings.daas_policy_enforce = True
    assert post(client, workbook_bytes(fixture_workbook())).status_code == 200


def test_filename_with_traversal_is_still_rejected_for_xlsx(client: TestClient, settings: Settings):
    r = post(client, workbook_bytes(fixture_workbook()), name="../../x.xlsx")
    assert r.status_code == 400
    assert_clean(settings)


# ---------------------------------------------------------------- round 3: limity, skan DTD, timeout, defusedxml
MIB = 1024 * 1024


def sheet_with_late_doctype(pad: int, marker: bytes = b"<!DOCTYPE x [<!ENTITY a 'b'>]>") -> bytes:
    """Poprawny XML, w ktorym DTD pojawia sie dopiero po `pad` bajtach (dalej niz dawny prolog 8 KiB)."""
    return (b'<?xml version="1.0"?><worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
            b"<!--" + b"x" * pad + b"-->" + marker + b"<sheetData/></worksheet>")


def test_unpacked_limit_is_48_mib():
    assert xi.MAX_XLSX_UNCOMPRESSED_BYTES == 48 * MIB


def test_unpacked_size_just_over_the_limit_is_rejected(tmp_path: Path, monkeypatch: pytest.MonkeyPatch):
    monkeypatch.setattr(xi, "RATIO_CHECK_FROM_BYTES", 10**12)  # sam wspolczynnik nie ma tu decydowac
    ok = tmp_path / "ok.xlsx"
    ok.write_bytes(zeros_zip(megabytes=44))
    xi.validate_container(ok)  # workbook (kilka KB) + 44 MiB < 48 MiB

    big = tmp_path / "big.xlsx"
    big.write_bytes(zeros_zip(megabytes=49))
    with pytest.raises(xi.UploadRejected) as exc:
        xi.validate_container(big)
    assert exc.value.reason == ReasonCode.XLSX_UNPACKED_TOO_LARGE and exc.value.http_status == 413


def with_part(name: str, body: bytes) -> bytes:
    """Skoroszyt z podmieniona (albo dodana, gdy openpyxl jej nie zapisal) czescia `name`."""
    raw = rewrite_zip(workbook_bytes(fixture_workbook()), lambda n, d: None if n == name else (n, d))
    return add_entry(raw, name, body)


def shared_strings_zip(body: bytes) -> bytes:
    return with_part("xl/sharedStrings.xml", body)



def test_shared_strings_entry_count_limit(tmp_path: Path):
    assert xi.MAX_SHARED_STRINGS == 200_000
    item = b"<si><t>a</t></si>"
    head = b'<?xml version="1.0"?><sst xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
    at_limit = tmp_path / "at.xlsx"
    at_limit.write_bytes(shared_strings_zip(head + item * xi.MAX_SHARED_STRINGS + b"</sst>"))
    xi.validate_container(at_limit)

    over = tmp_path / "over.xlsx"
    over.write_bytes(shared_strings_zip(head + item * (xi.MAX_SHARED_STRINGS + 1) + b"</sst>"))
    with pytest.raises(xi.UploadRejected) as exc:
        xi.validate_container(over)
    assert exc.value.reason == ReasonCode.XLSX_UNPACKED_TOO_LARGE


def test_shared_strings_prefixed_tags_are_counted_too(tmp_path: Path, monkeypatch: pytest.MonkeyPatch):
    monkeypatch.setattr(xi, "MAX_SHARED_STRINGS", 10)
    body = (b'<?xml version="1.0"?><x:sst xmlns:x="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
            + b"<x:si><x:t>a</x:t></x:si>" * 11 + b"</x:sst>")
    p = tmp_path / "p.xlsx"
    p.write_bytes(shared_strings_zip(body))
    with pytest.raises(xi.UploadRejected) as exc:
        xi.validate_container(p)
    assert exc.value.reason == ReasonCode.XLSX_UNPACKED_TOO_LARGE


def test_shared_strings_count_is_exact_across_chunk_boundaries(tmp_path: Path, monkeypatch: pytest.MonkeyPatch):
    monkeypatch.setattr(xi, "_SCAN_CHUNK_BYTES", 7)  # kazdy tag <si> trafia na granice porcji
    monkeypatch.setattr(xi, "MAX_SHARED_STRINGS", 50)
    head = b'<?xml version="1.0"?><sst xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
    exact = tmp_path / "exact.xlsx"
    exact.write_bytes(shared_strings_zip(head + b"<si><t>a</t></si>" * 50 + b"</sst>"))
    xi.validate_container(exact)  # 50 wpisow = limit: bez podwojnego liczenia na nakladkach
    over = tmp_path / "over.xlsx"
    over.write_bytes(shared_strings_zip(head + b"<si><t>a</t></si>" * 51 + b"</sst>"))
    with pytest.raises(xi.UploadRejected):
        xi.validate_container(over)


def test_shared_strings_size_limit(tmp_path: Path):
    assert xi.MAX_SHARED_STRINGS_BYTES == 24 * MIB
    head = b'<?xml version="1.0"?><sst xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><si><t>'
    p = tmp_path / "s.xlsx"
    p.write_bytes(shared_strings_zip(head + b"a" * (25 * MIB) + b"</t></si></sst>"))
    with pytest.raises(xi.UploadRejected) as exc:
        xi.validate_container(p)
    assert exc.value.reason == ReasonCode.XLSX_UNPACKED_TOO_LARGE and exc.value.http_status == 413


@pytest.mark.parametrize("part", ["docProps/core.xml", "xl/worksheets/sheet1.xml", "xl/sharedStrings.xml",
                                  "xl/_rels/workbook.xml.rels"])
@pytest.mark.parametrize("marker", [b"<!DOCTYPE x>", b"<!ENTITY a 'b'>"])
def test_doctype_far_beyond_the_old_prolog_is_rejected(client: TestClient, settings: Settings, part: str,
                                                      marker: bytes):
    raw = with_part(part, sheet_with_late_doctype(2 * MIB, marker))
    r = post(client, raw)
    assert r.status_code == 415
    assert audit(settings)[-1]["reason_code"] == "xlsx_invalid"
    assert_clean(settings)


@pytest.mark.parametrize("offset", range(0, 12))
def test_doctype_straddling_a_chunk_boundary_is_found(monkeypatch: pytest.MonkeyPatch, offset: int):
    monkeypatch.setattr(xi, "_SCAN_CHUNK_BYTES", 16)
    data = b"<a>" + b" " * (16 - 3 + offset) + b"<!DOCTYPE x>" + b"</a>"
    with pytest.raises(xi.UploadRejected):
        xi._scan_xml_part(io.BytesIO(data), count_strings=False, cancel=None)


def test_clean_large_xml_part_passes_the_scan():
    data = b"<a>" + b"<b>&lt;!DOCTYPE</b>" * 50_000 + b"</a>"  # zakodowane DOCTYPE w tekscie nie jest DTD
    xi._scan_xml_part(io.BytesIO(data), count_strings=False, cancel=None)


@pytest.mark.filterwarnings("ignore:Duplicate name")
def test_duplicate_zip_entry_names_are_all_scanned(client: TestClient):
    raw = workbook_bytes(fixture_workbook())
    out = io.BytesIO()
    evil = sheet_with_late_doctype(10)
    with zipfile.ZipFile(io.BytesIO(raw)) as zin, zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zout:
        zout.writestr("xl/worksheets/sheet1.xml", evil)  # pierwszy wpis: zlosliwy
        for info in zin.infolist():
            zout.writestr(info, zin.read(info.filename))  # drugi (ostatni) wpis o tej samej nazwie: czysty
    assert post(client, out.getvalue()).status_code == 415


def test_slow_parse_gives_422_xlsx_timeout(client: TestClient, settings: Settings, monkeypatch: pytest.MonkeyPatch):
    import time

    monkeypatch.setattr(xi, "XLSX_PARSE_TIMEOUT_SEC", 0.2)
    seen: list[object] = []

    def slow(ws, cancel=None):
        seen.append(cancel)
        for _ in range(100):  # ~5 s, ale konczy sie zaraz po ustawieniu flagi
            if cancel is not None and cancel.is_set():
                return b"", 0
            time.sleep(0.05)
        return b"", 0

    monkeypatch.setattr(xi, "_sheet_to_csv", slow)
    started = time.monotonic()
    r = post(client, workbook_bytes(fixture_workbook()))
    assert time.monotonic() - started < 3
    assert r.status_code == 422
    assert audit(settings)[-1]["reason_code"] == "xlsx_timeout"
    assert_clean(settings)
    assert seen and seen[0] is not None and seen[0].is_set()  # watek dostal sygnal przerwania


def test_timeout_reason_is_in_the_v1_error_code(monkeypatch: pytest.MonkeyPatch, tmp_path: Path):
    import time

    monkeypatch.setattr(xi, "XLSX_PARSE_TIMEOUT_SEC", 0.1)
    monkeypatch.setattr(xi, "validate_container", lambda path, cancel=None: time.sleep(0.5))
    p = tmp_path / "a.xlsx"
    p.write_bytes(b"x")
    with pytest.raises(xi.UploadRejected) as exc:
        xi.xlsx_to_csv_bytes(p)
    assert exc.value.reason == "xlsx_timeout" and exc.value.http_status == 422


def test_row_loop_stops_when_cancelled():
    import threading

    class Rows:
        def iter_rows(self, values_only=True):
            while True:
                yield ("a", "b")

    cancel = threading.Event()
    cancel.set()
    with pytest.raises(xi._Cancelled):
        xi._sheet_to_csv(Rows(), cancel)


def test_fast_conversion_returns_normally_within_the_timeout():
    wb = numbered_workbook(50)
    buf = io.BytesIO()
    wb.save(buf)
    fd, name = __import__("tempfile").mkstemp(suffix=".xlsx")
    __import__("os").close(fd)  # otwarty deskryptor blokowalby unlink na Windows
    p = Path(name)
    try:
        p.write_bytes(buf.getvalue())
        data, rows = xi.xlsx_to_csv_bytes(p)
    finally:
        p.unlink(missing_ok=True)
    assert rows == 50 and data.startswith(b"")


def test_openpyxl_parses_through_defusedxml():
    import openpyxl.xml as ox

    assert ox.DEFUSEDXML is True


def test_timeout_waits_for_the_worker_to_release_the_file(monkeypatch: pytest.MonkeyPatch):
    """Po timeoucie `run_bounded` daje watkowi chwile na zamkniecie pliku (Windows: inaczej unlink = WinError 32)."""
    import threading
    import time

    monkeypatch.setattr(xi, "XLSX_PARSE_TIMEOUT_SEC", 0.1)
    monkeypatch.setattr(xi, "XLSX_CANCEL_GRACE_SEC", 2.0)
    finished = threading.Event()

    def worker(cancel: threading.Event):
        try:
            for _ in range(200):
                if cancel.is_set():
                    return "cancelled"
                time.sleep(0.02)
            return "done"
        finally:
            finished.set()

    started = time.monotonic()
    with pytest.raises(xi.UploadRejected) as info:
        xi.run_bounded(worker)
    assert info.value.reason == "xlsx_timeout"
    assert finished.is_set()  # watek zakonczyl sie ZANIM wywolujacy dostal 422
    assert time.monotonic() - started < 1.5


def test_timeout_does_not_hang_on_a_worker_that_ignores_cancel(monkeypatch: pytest.MonkeyPatch):
    import threading
    import time

    monkeypatch.setattr(xi, "XLSX_PARSE_TIMEOUT_SEC", 0.05)
    monkeypatch.setattr(xi, "XLSX_CANCEL_GRACE_SEC", 0.1)
    release = threading.Event()

    def stubborn(cancel: threading.Event):
        release.wait(5)
        return "late"

    started = time.monotonic()
    try:
        with pytest.raises(xi.UploadRejected):
            xi.run_bounded(stubborn)
        assert time.monotonic() - started < 1.0  # laska jest ograniczona, odpowiedz nie czeka na uparty watek
    finally:
        release.set()


def test_remove_quietly_retries_a_locked_file_and_gives_up_quietly(tmp_path: Path, monkeypatch: pytest.MonkeyPatch):
    from app.core import upload_guard as ug

    target = tmp_path / "locked.tmp"
    target.write_bytes(b"x")
    denials: list[int] = []
    pauses: list[float] = []
    real_unlink = Path.unlink

    def flaky_unlink(self: Path, missing_ok: bool = False):
        if self == target and len(denials) < 2:
            denials.append(1)
            raise PermissionError(32, "sharing violation")
        return real_unlink(self, missing_ok=missing_ok)

    monkeypatch.setattr(Path, "unlink", flaky_unlink)
    assert ug.remove_quietly(target, sleep=pauses.append) is True
    assert not target.exists() and len(denials) == 2 and pauses == list(ug._REMOVE_BACKOFF_SEC[:2])

    always_locked = tmp_path / "forever.tmp"
    always_locked.write_bytes(b"x")

    def locked_unlink(self: Path, missing_ok: bool = False):
        if self == always_locked:
            raise PermissionError(32, "sharing violation")
        return real_unlink(self, missing_ok=missing_ok)

    monkeypatch.setattr(Path, "unlink", locked_unlink)
    pauses.clear()
    assert ug.remove_quietly(always_locked, sleep=pauses.append) is False  # bez wyjatku, plik zostaje dla sweepa
    assert len(pauses) == len(ug._REMOVE_BACKOFF_SEC) and always_locked.exists()
    monkeypatch.undo()
    assert ug.remove_quietly(tmp_path / "missing.tmp") is True
