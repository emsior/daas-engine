"""API v1: organizacje, role, czlonkowie, zaproszenia, klucze API, uploady."""
from __future__ import annotations

import csv
import io
import json
import re
from datetime import timedelta
from pathlib import Path

import openpyxl
import pytest
from fastapi.testclient import TestClient
from sqlalchemy import select

from app.control import models
from app.control import scopes as sc
from app.control.security import sha256_hex
from app.control.types import utcnow
from app.core.config import Settings
from tests.conftest import PROJECT_ROOT
from tests.v1_helpers import (
    API,
    get_rows,
    key_headers,
    login_via_magic_link,
    outbox,
    seed_key,
    seed_member,
    seed_org,
    seed_session,
    update_rows,
    user_headers,
    with_db,
)

CSV_BYTES = (PROJECT_ROOT / "data" / "fixtures" / "client_export_pl.csv").read_bytes()
PACKAGE_BYTES = (PROJECT_ROOT / "data" / "fixtures" / "uksc" / "host_compliant.json").read_bytes()


@pytest.fixture()
def org(cloud_client: TestClient, cloud_settings: Settings):
    return seed_org(cloud_settings)


def as_role(settings: Settings, org_id: str, role: str, email: str | None = None) -> dict[str, str]:
    """Naglowki uzytkownika z dana rola w organizacji."""
    email = email or f"{role}-{org_id[-6:].lower()}@example.test"
    return user_headers(settings, seed_member(settings, org_id, email, role))


def xlsx_bytes(rows: list[list[str]], title: str = "Dane") -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = title
    for row in rows:
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def csv_rows() -> list[list[str]]:
    return list(csv.reader(io.StringIO(CSV_BYTES.decode("utf-8-sig")), delimiter=";" if b";" in CSV_BYTES[:300] else ","))


# ================================================================ onboarding
def onboarding_body(**over) -> dict:
    return {"name": "Sklep Przykładowy", "dpa_accepted": True, "dpa_version": "v1", "locale": "pl", **over}


def test_create_org_makes_the_user_owner(cloud_client: TestClient, cloud_settings: Settings):
    hdr = user_headers(cloud_settings, seed_member(cloud_settings, seed_org(cloud_settings).org_id, "u@example.test",
                                                   "member"))
    r = cloud_client.post(f"{API}/orgs", json=onboarding_body(), headers=hdr)
    assert r.status_code == 201, r.text
    body = r.json()
    assert r.headers["location"] == f"/api/v1/orgs/{body['id']}"
    assert (body["plan"], body["retention_days"], body["locale"], body["dpa_version"], body["llm_narrative"]) == (
        "free", 90, "pl", "v1", False)
    assert body["dpa_accepted_at"] and body["slug"] == "sklep-przykladowy"
    me = cloud_client.get(f"{API}/me", headers=hdr).json()
    assert {(o["org_id"], o["role"]) for o in me["orgs"]} >= {(body["id"], "owner")}


def test_create_org_through_the_real_login_flow(cloud_client: TestClient, cloud_settings: Settings):
    login_via_magic_link(cloud_client, cloud_settings, "founder@example.test")
    csrf = cloud_client.cookies.get("daas_csrf")
    r = cloud_client.post(f"{API}/orgs", json=onboarding_body(locale="en"), headers={"X-CSRF-Token": csrf})
    assert r.status_code == 201 and r.json()["locale"] == "en"
    (session,) = get_rows(cloud_settings, models.Session)
    assert session.org_id == r.json()["id"]  # pierwsza organizacja staje sie aktywna


def test_create_org_is_audited(cloud_client: TestClient, cloud_settings: Settings, org):
    r = cloud_client.post(f"{API}/orgs", json=onboarding_body(), headers=user_headers(cloud_settings, org.user_id))
    (event,) = get_rows(cloud_settings, models.AuditEvent, models.AuditEvent.action == "org.created")
    assert event.org_id == r.json()["id"] and event.actor_id == org.user_id and event.meta == {"plan": "free"}


@pytest.mark.parametrize("body", [
    {"name": "X", "dpa_accepted": False},
    {"name": "X", "dpa_accepted": False, "dpa_version": "v1"},
])
def test_dpa_must_be_accepted(cloud_client: TestClient, cloud_settings: Settings, org, body):
    r = cloud_client.post(f"{API}/orgs", json=body, headers=user_headers(cloud_settings, org.user_id))
    assert r.status_code == 422 and r.json()["code"] == "dpa_required"


@pytest.mark.parametrize("body", [
    {"name": "X"},
    {"dpa_accepted": True},
    {"name": "", "dpa_accepted": True},
    {"name": "   ", "dpa_accepted": True},
    {"name": "N" * 201, "dpa_accepted": True},
    {"name": "X", "dpa_accepted": True, "dpa_version": "v2"},
    {"name": "X", "dpa_accepted": True, "locale": "de"},
    {"name": "X", "dpa_accepted": "yes please"},
])
def test_onboarding_validation(cloud_client: TestClient, cloud_settings: Settings, org, body):
    r = cloud_client.post(f"{API}/orgs", json=body, headers=user_headers(cloud_settings, org.user_id))
    assert r.status_code == 422 and r.headers["content-type"].startswith("application/problem+json")


def test_api_key_cannot_create_orgs(cloud_client: TestClient, org):
    r = cloud_client.post(f"{API}/orgs", json=onboarding_body(), headers=key_headers(org.api_key))
    assert r.status_code == 403 and r.json()["code"] == "user_session_required"


def test_anonymous_cannot_create_orgs(cloud_client: TestClient):
    assert cloud_client.post(f"{API}/orgs", json=onboarding_body()).status_code == 401


def test_create_org_needs_csrf_with_cookie(cloud_client: TestClient, cloud_settings: Settings, org):
    hdr = user_headers(cloud_settings, org.user_id)
    hdr.pop("X-CSRF-Token")
    r = cloud_client.post(f"{API}/orgs", json=onboarding_body(), headers=hdr)
    assert r.status_code == 403 and r.json()["code"] == "csrf_failed"


def test_create_org_is_rate_limited(cloud_client: TestClient, cloud_settings: Settings, org):
    hdr = user_headers(cloud_settings, org.user_id)
    codes = [cloud_client.post(f"{API}/orgs", json=onboarding_body(name=f"Firma {i}"), headers=hdr).status_code
             for i in range(7)]
    assert codes == [201] * 5 + [429, 429]


def test_same_name_gets_a_distinct_slug(cloud_client: TestClient, cloud_settings: Settings, org):
    hdr = user_headers(cloud_settings, org.user_id)
    slugs = [cloud_client.post(f"{API}/orgs", json=onboarding_body(name="Acme"), headers=hdr).json()["slug"]
             for _ in range(2)]
    assert slugs[0] != slugs[1] and slugs[1].startswith("acme-")


def test_selfhost_has_a_single_organization(anon_client: TestClient, settings: Settings):
    seeded_default = seed_member(settings, "default", "o@example.test", "owner")
    r = anon_client.post(f"{API}/orgs", json=onboarding_body(), headers=user_headers(settings, seeded_default))
    assert r.status_code == 403 and r.json()["code"] == "orgs_disabled_in_selfhost"


# ================================================================ odczyt i zmiana
def test_get_org(cloud_client: TestClient, org):
    r = cloud_client.get(f"{API}/orgs/{org.org_id}", headers=key_headers(org.api_key))
    assert r.status_code == 200 and r.json()["id"] == org.org_id and r.json()["name"] == "Acme"


def test_foreign_org_is_404_not_403(cloud_client: TestClient, cloud_settings: Settings, org):
    other = seed_org(cloud_settings, name="Other", owner_email="other@example.test")
    assert cloud_client.get(f"{API}/orgs/{other.org_id}", headers=key_headers(org.api_key)).status_code == 404
    assert cloud_client.get(f"{API}/orgs/{other.org_id}", headers=user_headers(cloud_settings, org.user_id)
                            ).status_code == 404
    assert cloud_client.get(f"{API}/orgs/NOSUCHORG", headers=key_headers(org.api_key)).status_code == 404


def test_patch_org_settings(cloud_client: TestClient, cloud_settings: Settings, org):
    r = cloud_client.patch(f"{API}/orgs/{org.org_id}", json={"retention_days": 30, "llm_narrative": True},
                           headers=key_headers(org.api_key))
    assert r.status_code == 200 and (r.json()["retention_days"], r.json()["llm_narrative"]) == (30, True)
    row = get_rows(cloud_settings, models.Org, models.Org.id == org.org_id)[0]
    assert (row.retention_days, row.llm_narrative) == (30, True)
    (event,) = get_rows(cloud_settings, models.AuditEvent, models.AuditEvent.action == "org.updated")
    assert event.meta == {"retention_days": 30, "llm_narrative": True}


def test_patch_name_and_locale(cloud_client: TestClient, org):
    r = cloud_client.patch(f"{API}/orgs/{org.org_id}", json={"name": "Nowa nazwa", "locale": "en"},
                           headers=key_headers(org.api_key))
    assert r.json()["name"] == "Nowa nazwa" and r.json()["locale"] == "en"


def test_empty_patch_changes_nothing(cloud_client: TestClient, org):
    assert cloud_client.patch(f"{API}/orgs/{org.org_id}", json={}, headers=key_headers(org.api_key)).status_code == 200


@pytest.mark.parametrize("body", [{"retention_days": 0}, {"retention_days": 4000}, {"retention_days": "x"},
                                  {"llm_narrative": "maybe"}, {"name": ""}, {"locale": "xx"}])
def test_patch_validation(cloud_client: TestClient, org, body):
    assert cloud_client.patch(f"{API}/orgs/{org.org_id}", json=body, headers=key_headers(org.api_key)
                              ).status_code == 422


def test_patch_cannot_change_plan_or_id(cloud_client: TestClient, org):
    r = cloud_client.patch(f"{API}/orgs/{org.org_id}", json={"plan": "pro", "id": "x", "slug": "y"},
                           headers=key_headers(org.api_key))
    assert r.json()["plan"] == "free" and r.json()["id"] == org.org_id and r.json()["slug"] == "acme"


def test_patch_needs_manage_scope(cloud_client: TestClient, cloud_settings: Settings, org):
    _, raw = seed_key(cloud_settings, org.org_id, scopes=[sc.ORG_READ])
    r = cloud_client.patch(f"{API}/orgs/{org.org_id}", json={"retention_days": 7}, headers=key_headers(raw))
    assert r.status_code == 403 and r.json()["required_scope"] == sc.ORG_MANAGE


# --- zakresy i audyt odmowy (zob. test_auth_v1)
def test_missing_scope_is_403_with_required_scope(cloud_client: TestClient, cloud_settings: Settings, org):
    _, raw = seed_key(cloud_settings, org.org_id, scopes=[sc.ORG_READ])
    r = cloud_client.get(f"{API}/orgs/{org.org_id}/api-keys", headers=key_headers(raw))
    assert r.status_code == 403 and r.headers["content-type"].startswith("application/problem+json")
    assert r.json()["code"] == "insufficient_scope" and r.json()["required_scope"] == sc.KEYS_MANAGE


def test_scope_denial_is_audited_in_the_org(cloud_client: TestClient, cloud_settings: Settings, org):
    _, raw = seed_key(cloud_settings, org.org_id, scopes=[sc.ORG_READ])
    cloud_client.get(f"{API}/orgs/{org.org_id}/api-keys", headers=key_headers(raw))
    (event,) = get_rows(cloud_settings, models.AuditEvent, models.AuditEvent.action == "auth_denied")
    assert event.org_id == org.org_id and event.actor_type == "key"
    assert event.meta == {"reason": "insufficient_scope", "scope": sc.KEYS_MANAGE}


# ================================================================ usuniecie
def test_owner_deletes_org_with_data_and_records(cloud_client: TestClient, cloud_settings: Settings, org):
    cloud_client.post(f"{API}/orgs/{org.org_id}/uploads", files={"file": ("a.csv", CSV_BYTES, "text/csv")},
                      headers=key_headers(org.api_key))
    tenant = cloud_settings.tenants_root / org.org_id
    assert tenant.is_dir()
    r = cloud_client.delete(f"{API}/orgs/{org.org_id}", headers=user_headers(cloud_settings, org.user_id))
    assert r.status_code == 204 and not tenant.exists()
    assert not get_rows(cloud_settings, models.Org, models.Org.id == org.org_id)
    assert not get_rows(cloud_settings, models.ApiKey) and not get_rows(cloud_settings, models.Upload)
    assert cloud_client.get(f"{API}/orgs/{org.org_id}", headers=key_headers(org.api_key)).status_code == 401
    (event,) = get_rows(cloud_settings, models.AuditEvent, models.AuditEvent.action == "org.deleted")
    assert event.org_id is None and event.target == f"org:{org.org_id}"


@pytest.mark.parametrize("role", ["admin", "member"])
def test_only_owner_can_delete_org(cloud_client: TestClient, cloud_settings: Settings, org, role):
    r = cloud_client.delete(f"{API}/orgs/{org.org_id}", headers=as_role(cloud_settings, org.org_id, role))
    assert r.status_code == 403 and r.json()["required_scope"] == sc.ORG_DELETE


def test_api_key_cannot_delete_org(cloud_client: TestClient, org):
    assert cloud_client.delete(f"{API}/orgs/{org.org_id}", headers=key_headers(org.api_key)).status_code == 403


def test_delete_is_refused_while_a_job_is_active(cloud_client: TestClient, cloud_settings: Settings, org):
    async def add(db):
        db.add(models.Job(org_id=org.org_id, pipeline="cs2_demo", status="running"))

    with_db(cloud_settings, add)
    r = cloud_client.delete(f"{API}/orgs/{org.org_id}", headers=user_headers(cloud_settings, org.user_id))
    assert r.status_code == 409 and r.json()["code"] == "org_busy"
    assert get_rows(cloud_settings, models.Org, models.Org.id == org.org_id)


def test_selfhost_org_cannot_be_deleted(anon_client: TestClient, settings: Settings):
    hdr = user_headers(settings, seed_member(settings, "default", "o@example.test", "owner"))
    r = anon_client.delete(f"{API}/orgs/default", headers=hdr)
    assert r.status_code == 403 and r.json()["code"] == "org_protected"


def test_delete_after_delete_is_404(cloud_client: TestClient, cloud_settings: Settings, org):
    hdr = user_headers(cloud_settings, org.user_id)
    cloud_client.delete(f"{API}/orgs/{org.org_id}", headers=hdr)
    assert cloud_client.delete(f"{API}/orgs/{org.org_id}", headers=hdr).status_code == 404


# ================================================================ czlonkowie
def test_list_members_for_every_role(cloud_client: TestClient, cloud_settings: Settings, org):
    seed_member(cloud_settings, org.org_id, "adm@example.test", "admin")
    for role in ("owner", "admin", "member"):
        hdr = as_role(cloud_settings, org.org_id, role, email=f"{role}@example.test")
        rows = cloud_client.get(f"{API}/orgs/{org.org_id}/members", headers=hdr).json()
        assert {r["email"] for r in rows} >= {org.email, "adm@example.test"}
        assert set(rows[0]) == {"user_id", "email", "name", "role", "created_at"}


def test_non_member_gets_403_not_404_on_members(cloud_client: TestClient, cloud_settings: Settings, org):
    other = seed_org(cloud_settings, name="Other", owner_email="other@example.test")
    assert cloud_client.get(f"{API}/orgs/{org.org_id}/members", headers=user_headers(cloud_settings, other.user_id)
                            ).status_code == 403
    assert cloud_client.get(f"{API}/orgs/{org.org_id}/members", headers=key_headers(other.api_key)).status_code == 403
    assert cloud_client.get(f"{API}/orgs/NOSUCH/members", headers=key_headers(other.api_key)).status_code == 403


def test_add_existing_user(cloud_client: TestClient, cloud_settings: Settings, org):
    seed_member(cloud_settings, seed_org(cloud_settings, name="B", owner_email="b@example.test").org_id,
                "ola@example.test", "member")
    r = cloud_client.post(f"{API}/orgs/{org.org_id}/members", json={"email": "Ola@Example.test", "role": "admin"},
                          headers=key_headers(org.api_key))
    assert r.status_code == 201 and r.json()["email"] == "ola@example.test" and r.json()["role"] == "admin"
    (event,) = get_rows(cloud_settings, models.AuditEvent, models.AuditEvent.action == "member.added")
    assert event.meta == {"role": "admin"}


def test_add_unknown_user_is_404_with_hint(cloud_client: TestClient, org):
    r = cloud_client.post(f"{API}/orgs/{org.org_id}/members", json={"email": "ghost@example.test"},
                          headers=key_headers(org.api_key))
    assert r.status_code == 404 and r.json()["code"] == "user_not_found" and "invitation" in r.json()["detail"]


def test_add_duplicate_member_is_409(cloud_client: TestClient, org):
    r = cloud_client.post(f"{API}/orgs/{org.org_id}/members", json={"email": org.email},
                          headers=key_headers(org.api_key))
    assert r.status_code == 409 and r.json()["code"] == "already_member"


@pytest.mark.parametrize("role", ["member"])
def test_plain_member_cannot_manage_members(cloud_client: TestClient, cloud_settings: Settings, org, role):
    hdr = as_role(cloud_settings, org.org_id, role)
    assert cloud_client.post(f"{API}/orgs/{org.org_id}/members", json={"email": org.email}, headers=hdr
                             ).status_code == 403
    assert cloud_client.post(f"{API}/orgs/{org.org_id}/invitations", json={"email": "n@example.test"}, headers=hdr
                             ).status_code == 403


def test_admin_cannot_grant_owner_but_owner_can(cloud_client: TestClient, cloud_settings: Settings, org):
    seed_member(cloud_settings, seed_org(cloud_settings, name="B", owner_email="b@example.test").org_id,
                "new@example.test", "member")
    admin = as_role(cloud_settings, org.org_id, "admin")
    r = cloud_client.post(f"{API}/orgs/{org.org_id}/members", json={"email": "new@example.test", "role": "owner"},
                          headers=admin)
    assert r.status_code == 403
    owner = user_headers(cloud_settings, org.user_id)
    r = cloud_client.post(f"{API}/orgs/{org.org_id}/members", json={"email": "new@example.test", "role": "owner"},
                          headers=owner)
    assert r.status_code == 201 and r.json()["role"] == "owner"


def test_api_key_cannot_grant_owner(cloud_client: TestClient, org):
    r = cloud_client.post(f"{API}/orgs/{org.org_id}/invitations", json={"email": "n@example.test", "role": "owner"},
                          headers=key_headers(org.api_key))
    assert r.status_code == 403


@pytest.mark.parametrize("email", ["", "no-at", "a@b", None])
def test_member_email_validation(cloud_client: TestClient, org, email):
    r = cloud_client.post(f"{API}/orgs/{org.org_id}/members", json={"email": email}, headers=key_headers(org.api_key))
    assert r.status_code == 422


def test_invalid_role_is_422(cloud_client: TestClient, org):
    r = cloud_client.post(f"{API}/orgs/{org.org_id}/invitations", json={"email": "n@example.test", "role": "root"},
                          headers=key_headers(org.api_key))
    assert r.status_code == 422


def test_invitation_is_stored_hashed_mailed_and_audited(cloud_client: TestClient, cloud_settings: Settings, org):
    r = cloud_client.post(f"{API}/orgs/{org.org_id}/invitations", json={"email": "New@Example.test", "role": "member"},
                          headers=key_headers(org.api_key))
    assert r.status_code == 201 and r.json()["email"] == "new@example.test" and r.json()["role"] == "member"
    assert timedelta(days=6) < cloud_settings_expiry(cloud_settings) <= timedelta(days=7)
    (mail,) = outbox(cloud_settings, "new@example.test")
    assert "Acme" in mail["subject"] and mail["link"].endswith("/login")
    (event,) = get_rows(cloud_settings, models.AuditEvent, models.AuditEvent.action == "member.invited")
    assert event.meta["role"] == "member" and "new@example.test" not in json.dumps(event.meta)


def cloud_settings_expiry(settings: Settings) -> timedelta:
    (inv,) = get_rows(settings, models.Invitation)
    return inv.expires_at - utcnow()


def test_repeating_an_invitation_refreshes_it(cloud_client: TestClient, cloud_settings: Settings, org):
    ids = [cloud_client.post(f"{API}/orgs/{org.org_id}/invitations", json={"email": "n@example.test", "role": role},
                             headers=key_headers(org.api_key)).json()["id"] for role in ("member", "admin")]
    assert ids[0] == ids[1]
    (inv,) = get_rows(cloud_settings, models.Invitation)
    assert inv.role == "admin"


def test_inviting_an_existing_member_is_409(cloud_client: TestClient, org):
    r = cloud_client.post(f"{API}/orgs/{org.org_id}/invitations", json={"email": org.email},
                          headers=key_headers(org.api_key))
    assert r.status_code == 409


def test_invited_person_joins_on_first_login(cloud_client: TestClient, cloud_settings: Settings, org):
    cloud_client.post(f"{API}/orgs/{org.org_id}/invitations", json={"email": "joiner@example.test", "role": "admin"},
                      headers=key_headers(org.api_key))
    body = login_via_magic_link(cloud_client, cloud_settings, "joiner@example.test")
    assert [(o["org_id"], o["role"]) for o in body["orgs"]] == [(org.org_id, "admin")]


def test_invitation_mail_failure_does_not_fail_the_request(cloud_client: TestClient, org,
                                                           monkeypatch: pytest.MonkeyPatch):
    def boom(_mail):
        raise OSError("smtp down")

    monkeypatch.setattr(cloud_client.app.state.mailer, "send", boom)
    r = cloud_client.post(f"{API}/orgs/{org.org_id}/invitations", json={"email": "n@example.test"},
                          headers=key_headers(org.api_key))
    assert r.status_code == 201


# ================================================================ klucze API
def keys_url(org_id: str) -> str:
    return f"{API}/orgs/{org_id}/api-keys"


def test_create_key_shows_it_once(cloud_client: TestClient, cloud_settings: Settings, org):
    r = cloud_client.post(keys_url(org.org_id), json={"name": "n8n", "scopes": ["runs:read", "reports:read"]},
                          headers=key_headers(org.api_key))
    assert r.status_code == 201
    body = r.json()
    raw = body["key"]
    assert raw.startswith("daas_") and body["prefix"] == raw[:13] and body["scopes"] == ["reports:read", "runs:read"]
    assert body["last_used_at"] is None and body["revoked_at"] is None and body["expires_at"] is None
    listing = cloud_client.get(keys_url(org.org_id), headers=key_headers(org.api_key))
    assert raw not in listing.text and "key" not in listing.json()[0]
    row = [k for k in get_rows(cloud_settings, models.ApiKey) if k.id == body["id"]][0]
    assert row.key_hash == sha256_hex(raw) and raw not in str(row.__dict__)


def test_created_key_works_with_exactly_its_scopes(cloud_client: TestClient, org):
    raw = cloud_client.post(keys_url(org.org_id), json={"name": "ro", "scopes": ["runs:read"]},
                            headers=key_headers(org.api_key)).json()["key"]
    me = cloud_client.get(f"{API}/me", headers=key_headers(raw)).json()
    assert me["principal"]["scopes"] == ["runs:read"] and me["principal"]["org_id"] == org.org_id
    assert cloud_client.get(keys_url(org.org_id), headers=key_headers(raw)).status_code == 403


def test_key_expiry_is_set(cloud_client: TestClient, cloud_settings: Settings, org):
    body = cloud_client.post(keys_url(org.org_id), json={"name": "t", "scopes": ["runs:read"], "expires_in_days": 30},
                             headers=key_headers(org.api_key)).json()
    row = [k for k in get_rows(cloud_settings, models.ApiKey) if k.id == body["id"]][0]
    assert timedelta(days=29) < row.expires_at - utcnow() <= timedelta(days=30)


@pytest.mark.parametrize("body", [
    {"name": "x", "scopes": []}, {"name": "x", "scopes": ["root:all"]}, {"name": "", "scopes": ["runs:read"]},
    {"name": "  ", "scopes": ["runs:read"]}, {"scopes": ["runs:read"]}, {"name": "x"},
    {"name": "x", "scopes": ["runs:read"], "expires_in_days": 0},
    {"name": "x", "scopes": ["runs:read"], "expires_in_days": 99999},
    {"name": "x", "scopes": ["org:delete"]},
])
def test_key_validation(cloud_client: TestClient, org, body):
    assert cloud_client.post(keys_url(org.org_id), json=body, headers=key_headers(org.api_key)).status_code == 422


def test_scope_escalation_is_refused(cloud_client: TestClient, cloud_settings: Settings, org):
    _, narrow = seed_key(cloud_settings, org.org_id, scopes=[sc.KEYS_MANAGE, sc.RUNS_READ])
    r = cloud_client.post(keys_url(org.org_id), json={"name": "x", "scopes": ["runs:read", "uploads:write"]},
                          headers=key_headers(narrow))
    assert r.status_code == 403 and r.json()["code"] == "scope_escalation" and r.json()["scopes"] == ["uploads:write"]
    ok = cloud_client.post(keys_url(org.org_id), json={"name": "x", "scopes": ["runs:read"]},
                           headers=key_headers(narrow))
    assert ok.status_code == 201


@pytest.mark.parametrize(("role", "status"), [("owner", 201), ("admin", 201), ("member", 403)])
def test_key_management_by_role(cloud_client: TestClient, cloud_settings: Settings, org, role, status):
    hdr = as_role(cloud_settings, org.org_id, role, email=f"{role}@example.test")
    r = cloud_client.post(keys_url(org.org_id), json={"name": "x", "scopes": ["runs:read"]}, headers=hdr)
    assert r.status_code == status
    assert cloud_client.get(keys_url(org.org_id), headers=hdr).status_code == (200 if status == 201 else 403)


def test_key_creation_and_revocation_are_audited_without_the_key(cloud_client: TestClient, cloud_settings: Settings,
                                                                 org):
    made = cloud_client.post(keys_url(org.org_id), json={"name": "x", "scopes": ["runs:read"]},
                             headers=key_headers(org.api_key)).json()
    cloud_client.delete(f"{keys_url(org.org_id)}/{made['id']}", headers=key_headers(org.api_key))
    created = get_rows(cloud_settings, models.AuditEvent, models.AuditEvent.action == "key.created")[0]
    revoked = get_rows(cloud_settings, models.AuditEvent, models.AuditEvent.action == "key.revoked")[0]
    assert created.meta == {"key_prefix": made["key"][5:13], "scopes": ["runs:read"]}
    assert revoked.meta == {"key_prefix": made["key"][5:13]} and made["key"] not in json.dumps(created.meta)


def test_revoke_stops_the_key_at_once(cloud_client: TestClient, org):
    made = cloud_client.post(keys_url(org.org_id), json={"name": "x", "scopes": ["runs:read"]},
                             headers=key_headers(org.api_key)).json()
    assert cloud_client.get(f"{API}/me", headers=key_headers(made["key"])).status_code == 200
    assert cloud_client.delete(f"{keys_url(org.org_id)}/{made['id']}", headers=key_headers(org.api_key)
                               ).status_code == 204
    assert cloud_client.get(f"{API}/me", headers=key_headers(made["key"])).status_code == 401
    listing = cloud_client.get(keys_url(org.org_id), headers=key_headers(org.api_key)).json()
    assert [k["revoked_at"] is not None for k in listing if k["id"] == made["id"]] == [True]


def test_revoking_twice_is_idempotent(cloud_client: TestClient, cloud_settings: Settings, org):
    made = cloud_client.post(keys_url(org.org_id), json={"name": "x", "scopes": ["runs:read"]},
                             headers=key_headers(org.api_key)).json()
    for _ in range(2):
        assert cloud_client.delete(f"{keys_url(org.org_id)}/{made['id']}", headers=key_headers(org.api_key)
                                   ).status_code == 204
    assert len(get_rows(cloud_settings, models.AuditEvent, models.AuditEvent.action == "key.revoked")) == 1


def test_rotation_flow(cloud_client: TestClient, org):
    old = cloud_client.post(keys_url(org.org_id), json={"name": "v1", "scopes": ["runs:read"]},
                            headers=key_headers(org.api_key)).json()
    new = cloud_client.post(keys_url(org.org_id), json={"name": "v2", "scopes": ["runs:read"]},
                            headers=key_headers(org.api_key)).json()
    cloud_client.delete(f"{keys_url(org.org_id)}/{old['id']}", headers=key_headers(org.api_key))
    assert cloud_client.get(f"{API}/me", headers=key_headers(old["key"])).status_code == 401
    assert cloud_client.get(f"{API}/me", headers=key_headers(new["key"])).status_code == 200


def test_cannot_revoke_a_key_of_another_org(cloud_client: TestClient, cloud_settings: Settings, org):
    other = seed_org(cloud_settings, name="Other", owner_email="other@example.test")
    r = cloud_client.delete(f"{keys_url(org.org_id)}/{other.key_id}", headers=key_headers(org.api_key))
    assert r.status_code == 404
    assert cloud_client.get(f"{API}/me", headers=key_headers(other.api_key)).status_code == 200
    assert cloud_client.delete(f"{keys_url(org.org_id)}/NOSUCHKEY", headers=key_headers(org.api_key)
                               ).status_code == 404


def test_list_shows_newest_first(cloud_client: TestClient, org):
    for name in ("first", "second"):
        cloud_client.post(keys_url(org.org_id), json={"name": name, "scopes": ["runs:read"]},
                          headers=key_headers(org.api_key))
    assert [k["name"] for k in cloud_client.get(keys_url(org.org_id), headers=key_headers(org.api_key)).json()][:2
        ] == ["second", "first"]


def test_active_key_limit(cloud_client: TestClient, cloud_settings: Settings, org):
    from app.api.v1.api_keys import MAX_ACTIVE_KEYS

    for _ in range(MAX_ACTIVE_KEYS - 1):  # klucz startowy zajmuje jedno miejsce
        seed_key(cloud_settings, org.org_id)
    r = cloud_client.post(keys_url(org.org_id), json={"name": "x", "scopes": ["runs:read"]},
                          headers=key_headers(org.api_key))
    assert r.status_code == 409 and r.json()["code"] == "key_limit_reached"


# ================================================================ uploady
def uploads_url(org_id: str) -> str:
    return f"{API}/orgs/{org_id}/uploads"


def test_csv_upload_returns_mapping_and_stores_file(cloud_client: TestClient, cloud_settings: Settings, org):
    r = cloud_client.post(uploads_url(org.org_id), files={"file": ("dane.csv", CSV_BYTES, "text/csv")},
                          headers=key_headers(org.api_key))
    assert r.status_code == 201, r.text
    body = r.json()
    assert body["kind"] == "csv" and body["suggested_pipeline"] == "ecommerce_demo" and body["ready"] is True
    assert body["rows"] > 0 and "revenue" in body["mapping"] and body["bytes"] == len(CSV_BYTES)
    assert body["sha256"] == sha256_hex(CSV_BYTES) and body["filename"] == "dane.csv"
    stored = cloud_settings.tenants_root / org.org_id / "uploads" / f"{body['id']}.csv"
    assert stored.is_file() and not list(stored.parent.glob("*.tmp"))
    (row,) = get_rows(cloud_settings, models.Upload)
    assert (row.org_id, row.kind, row.path, row.bytes) == (org.org_id, "csv", stored.name, len(CSV_BYTES))
    assert timedelta(days=89) < row.retention_until - utcnow() <= timedelta(days=90)


def test_retention_follows_the_org_setting(cloud_client: TestClient, cloud_settings: Settings, org):
    cloud_client.patch(f"{API}/orgs/{org.org_id}", json={"retention_days": 7}, headers=key_headers(org.api_key))
    cloud_client.post(uploads_url(org.org_id), files={"file": ("a.csv", CSV_BYTES, "text/csv")},
                      headers=key_headers(org.api_key))
    (row,) = get_rows(cloud_settings, models.Upload)
    assert timedelta(days=6) < row.retention_until - utcnow() <= timedelta(days=7)


def test_upload_counts_usage_and_is_audited(cloud_client: TestClient, cloud_settings: Settings, org):
    for _ in range(2):
        cloud_client.post(uploads_url(org.org_id), files={"file": ("a.csv", CSV_BYTES, "text/csv")},
                          headers=key_headers(org.api_key))
    (usage,) = get_rows(cloud_settings, models.UsageCounter)
    assert (usage.org_id, usage.period, usage.upload_bytes, usage.runs) == (
        org.org_id, utcnow().strftime("%Y-%m"), 2 * len(CSV_BYTES), 0)
    events = get_rows(cloud_settings, models.AuditEvent, models.AuditEvent.action == "upload.created")
    assert len(events) == 2 and set(events[0].meta) == {"upload_id", "kind", "bytes", "sha256"}
    assert "dane" not in json.dumps(events[0].meta)
    lines = (cloud_settings.tenants_root / org.org_id / "audit.jsonl").read_text(encoding="utf-8").splitlines()
    assert [json.loads(x).get("action") for x in lines].count("upload.created") == 2


def test_xlsx_upload(cloud_client: TestClient, cloud_settings: Settings, org):
    data = xlsx_bytes(csv_rows())
    r = cloud_client.post(uploads_url(org.org_id), files={"file": ("dane.xlsx", data, "application/octet-stream")},
                          headers=key_headers(org.api_key))
    assert r.status_code == 201, r.text
    body = r.json()
    assert body["kind"] == "xlsx" and body["ready"] is True and body["sha256"] == sha256_hex(data)
    stored = cloud_settings.tenants_root / org.org_id / "uploads" / f"{body['id']}.csv"
    assert stored.is_file() and b"PK\x03\x04" not in stored.read_bytes()[:4]  # na dysku jest CSV, nie skoroszyt


def test_xlsx_with_several_sheets_needs_a_sheet_name(cloud_client: TestClient, org):
    wb = openpyxl.Workbook()
    wb.active.title = "A"
    for row in csv_rows():
        wb.active.append(row)
    wb.create_sheet("B")
    buf = io.BytesIO()
    wb.save(buf)
    files = {"file": ("w.xlsx", buf.getvalue(), "application/octet-stream")}
    r = cloud_client.post(uploads_url(org.org_id), files=files, headers=key_headers(org.api_key))
    assert r.status_code == 422 and r.json()["code"] == "xlsx_multiple_sheets"
    r = cloud_client.post(uploads_url(org.org_id), files=files, data={"sheet": "A"}, headers=key_headers(org.api_key))
    assert r.status_code == 201


def test_macro_xlsx_name_is_refused(cloud_client: TestClient, org):
    r = cloud_client.post(uploads_url(org.org_id), files={"file": ("m.xlsm", b"x", "application/octet-stream")},
                          headers=key_headers(org.api_key))
    assert r.status_code == 415 and r.json()["code"] == "extension_not_allowed"


def test_fake_xlsx_is_refused(cloud_client: TestClient, cloud_settings: Settings, org):
    r = cloud_client.post(uploads_url(org.org_id), files={"file": ("f.xlsx", b"not a zip", "application/octet-stream")},
                          headers=key_headers(org.api_key))
    assert r.status_code == 415 and r.json()["code"] == "xlsx_invalid"
    assert not get_rows(cloud_settings, models.Upload)


def test_package_upload(cloud_client: TestClient, cloud_settings: Settings, org):
    r = cloud_client.post(uploads_url(org.org_id), files={"file": ("host.json", PACKAGE_BYTES, "application/json")},
                          headers=key_headers(org.api_key))
    assert r.status_code == 201, r.text
    body = r.json()
    assert body["kind"] == "package" and body["suggested_pipeline"] == "uksc_evidence" and body["rows"] is None
    assert (cloud_settings.tenants_root / org.org_id / "uploads" / f"{body['id']}.json").read_bytes() == PACKAGE_BYTES


@pytest.mark.parametrize("payload", [b"{}", b"[]", b"not json", b'{"schema_version": "9.9"}', b"", b"\xff\xfe\x00"])
def test_invalid_package_is_refused_without_echo(cloud_client: TestClient, cloud_settings: Settings, org, payload):
    r = cloud_client.post(uploads_url(org.org_id), files={"file": ("p.json", payload, "application/json")},
                          headers=key_headers(org.api_key))
    assert r.status_code == 422 and not get_rows(cloud_settings, models.Upload)
    assert "schema_version" not in r.json()["detail"]
    assert not list((cloud_settings.tenants_root / org.org_id / "uploads").iterdir())


def test_package_over_5_mib_is_refused(cloud_client: TestClient, org):
    r = cloud_client.post(uploads_url(org.org_id),
                          files={"file": ("p.json", b" " * (5 * 1024 * 1024 + 1), "application/json")},
                          headers=key_headers(org.api_key))
    assert r.status_code == 413 and r.json()["code"] == "file_too_large"


def test_csv_over_10_mib_is_refused(cloud_client: TestClient, org):
    big = b"date,revenue\n" + b"2026-01-05,100\n" * 800_000
    r = cloud_client.post(uploads_url(org.org_id), files={"file": ("big.csv", big, "text/csv")},
                          headers=key_headers(org.api_key))
    assert r.status_code == 413 and r.json()["code"] == "file_too_large"


@pytest.mark.parametrize(("name", "status"), [
    ("a.txt", 415), ("a.exe", 415), ("a", 415), ("../a.csv", 400), ("a/b.csv", 400), ("a\\b.csv", 400),
    ("x" * 130 + ".csv", 400), ("..\\..\\a.json", 400),
])
def test_filename_rules(cloud_client: TestClient, cloud_settings: Settings, org, name, status):
    r = cloud_client.post(uploads_url(org.org_id), files={"file": (name, CSV_BYTES, "text/csv")},
                          headers=key_headers(org.api_key))
    assert r.status_code == status and not get_rows(cloud_settings, models.Upload)


def test_wrong_csv_schema_is_refused(cloud_client: TestClient, org):
    r = cloud_client.post(uploads_url(org.org_id), files={"file": ("a.csv", b"foo,bar\n1,2\n", "text/csv")},
                          headers=key_headers(org.api_key))
    assert r.status_code == 422 and r.json()["code"] == "schema_mismatch"


def test_upload_without_file_is_422(cloud_client: TestClient, org):
    assert cloud_client.post(uploads_url(org.org_id), headers=key_headers(org.api_key)).status_code == 422


def test_refused_upload_is_audited_without_leftovers(cloud_client: TestClient, cloud_settings: Settings, org):
    cloud_client.post(uploads_url(org.org_id), files={"file": ("a.csv", b"foo,bar\n1,2\n", "text/csv")},
                      headers=key_headers(org.api_key))
    (event,) = get_rows(cloud_settings, models.AuditEvent, models.AuditEvent.action == "auth_denied")
    assert event.meta["reason"] == "schema_mismatch" and event.target == "upload"
    assert not list((cloud_settings.tenants_root / org.org_id / "uploads").iterdir())


def test_quota_is_enforced_per_org(cloud_factory):
    client, settings = cloud_factory(DAAS_UPLOAD_QUOTA_BYTES=len(CSV_BYTES) + 10)
    a = seed_org(settings, name="A", owner_email="a@example.test")
    b = seed_org(settings, name="B", owner_email="b@example.test")
    assert client.post(uploads_url(a.org_id), files={"file": ("a.csv", CSV_BYTES, "text/csv")},
                       headers=key_headers(a.api_key)).status_code == 201
    r = client.post(uploads_url(a.org_id), files={"file": ("a.csv", CSV_BYTES, "text/csv")},
                    headers=key_headers(a.api_key))
    assert r.status_code == 413 and r.json()["code"] == "quota_exceeded"
    assert client.post(uploads_url(b.org_id), files={"file": ("a.csv", CSV_BYTES, "text/csv")},
                       headers=key_headers(b.api_key)).status_code == 201  # kwota liczy sie osobno


def test_upload_rate_limit(cloud_factory):
    client, settings = cloud_factory(DAAS_RATE_LIMIT_UPLOAD_PER_MIN=2)
    a = seed_org(settings)
    codes = [client.post(uploads_url(a.org_id), files={"file": ("a.csv", CSV_BYTES, "text/csv")},
                         headers=key_headers(a.api_key)).status_code for _ in range(4)]
    assert codes == [201, 201, 429, 429]


def test_member_may_upload_but_a_key_without_scope_may_not(cloud_client: TestClient, cloud_settings: Settings, org):
    hdr = as_role(cloud_settings, org.org_id, "member")
    assert cloud_client.post(uploads_url(org.org_id), files={"file": ("a.csv", CSV_BYTES, "text/csv")},
                             headers=hdr).status_code == 201
    _, ro = seed_key(cloud_settings, org.org_id, scopes=[sc.RUNS_READ])
    r = cloud_client.post(uploads_url(org.org_id), files={"file": ("a.csv", CSV_BYTES, "text/csv")},
                          headers=key_headers(ro))
    assert r.status_code == 403 and r.json()["required_scope"] == sc.UPLOADS_WRITE


def test_upload_with_cookie_needs_csrf(cloud_client: TestClient, cloud_settings: Settings, org):
    hdr = user_headers(cloud_settings, org.user_id)
    hdr.pop("X-CSRF-Token")
    r = cloud_client.post(uploads_url(org.org_id), files={"file": ("a.csv", CSV_BYTES, "text/csv")}, headers=hdr)
    assert r.status_code == 403 and r.json()["code"] == "csrf_failed"


def test_each_org_gets_its_own_upload_directory(cloud_client: TestClient, cloud_settings: Settings, org):
    other = seed_org(cloud_settings, name="Other", owner_email="other@example.test")
    for o in (org, other):
        cloud_client.post(uploads_url(o.org_id), files={"file": ("a.csv", CSV_BYTES, "text/csv")},
                          headers=key_headers(o.api_key))
    dirs = {o.org_id: list((cloud_settings.tenants_root / o.org_id / "uploads").iterdir()) for o in (org, other)}
    assert all(len(v) == 1 for v in dirs.values()) and dirs[org.org_id] != dirs[other.org_id]


def test_upload_into_a_foreign_org_is_404(cloud_client: TestClient, cloud_settings: Settings, org):
    other = seed_org(cloud_settings, name="Other", owner_email="other@example.test")
    r = cloud_client.post(uploads_url(other.org_id), files={"file": ("a.csv", CSV_BYTES, "text/csv")},
                          headers=key_headers(org.api_key))
    assert r.status_code == 404 and not get_rows(cloud_settings, models.Upload)
    assert not (cloud_settings.tenants_root / other.org_id / "uploads").exists()


def test_filename_with_unicode_is_kept_for_display_only(cloud_client: TestClient, cloud_settings: Settings, org):
    r = cloud_client.post(uploads_url(org.org_id), files={"file": ("sprzedaż_październik.csv", CSV_BYTES, "text/csv")},
                          headers=key_headers(org.api_key))
    assert r.status_code == 201 and r.json()["filename"] == "sprzedaż_październik.csv"
    stored = list((cloud_settings.tenants_root / org.org_id / "uploads").iterdir())
    assert len(stored) == 1 and stored[0].name == f"{r.json()['id']}.csv"


# ================================================================ macierz rol
ROLE_MATRIX = [
    # (rola, metoda, sciezka-szablon, body, oczekiwany status)
    ("member", "GET", "/orgs/{o}", None, 200),
    ("member", "PATCH", "/orgs/{o}", {"retention_days": 10}, 403),
    ("member", "GET", "/orgs/{o}/members", None, 200),
    ("member", "GET", "/orgs/{o}/api-keys", None, 403),
    ("admin", "GET", "/orgs/{o}/api-keys", None, 200),
    ("admin", "PATCH", "/orgs/{o}", {"retention_days": 10}, 200),
    ("admin", "DELETE", "/orgs/{o}", None, 403),
    ("owner", "GET", "/orgs/{o}/api-keys", None, 200),
    ("owner", "PATCH", "/orgs/{o}", {"retention_days": 10}, 200),
]


@pytest.mark.parametrize(("role", "method", "path", "body", "status"), ROLE_MATRIX)
def test_role_matrix(cloud_client: TestClient, cloud_settings: Settings, org, role, method, path, body, status):
    hdr = as_role(cloud_settings, org.org_id, role, email=f"{role}@example.test")
    r = cloud_client.request(method, API + path.format(o=org.org_id), json=body, headers=hdr)
    assert r.status_code == status, r.text


def test_session_user_in_two_orgs_sees_each_with_the_right_role(cloud_client: TestClient, cloud_settings: Settings,
                                                                org):
    other = seed_org(cloud_settings, name="Other", owner_email="other@example.test")
    seed_member(cloud_settings, other.org_id, org.email, "member")
    hdr = user_headers(cloud_settings, org.user_id)
    assert cloud_client.get(f"{API}/orgs/{other.org_id}/api-keys", headers=hdr).status_code == 403  # tylko member
    assert cloud_client.get(f"{API}/orgs/{org.org_id}/api-keys", headers=hdr).status_code == 200  # owner
    roles = {o["org_id"]: o["role"] for o in cloud_client.get(f"{API}/me", headers=hdr).json()["orgs"]}
    assert roles == {org.org_id: "owner", other.org_id: "member"}


def test_removed_membership_takes_effect_immediately(cloud_client: TestClient, cloud_settings: Settings, org):
    hdr = as_role(cloud_settings, org.org_id, "member", email="leaver@example.test")
    assert cloud_client.get(f"{API}/orgs/{org.org_id}", headers=hdr).status_code == 200

    async def drop(db):
        for m in (await db.execute(select(models.Membership).where(models.Membership.role == "member"))).scalars():
            await db.delete(m)

    with_db(cloud_settings, drop)
    assert cloud_client.get(f"{API}/orgs/{org.org_id}", headers=hdr).status_code == 404


def test_session_without_org_header_unchanged(cloud_client: TestClient, cloud_settings: Settings, org):
    cookie, _ = seed_session(cloud_settings, org.user_id)
    cloud_client.cookies.set("daas_session", cookie)
    assert cloud_client.get(f"{API}/orgs/{org.org_id}").status_code == 200
    update_rows(cloud_settings, models.Session, {"expires_at": utcnow() - timedelta(seconds=1)})
    assert cloud_client.get(f"{API}/orgs/{org.org_id}").status_code == 401


def test_openapi_has_no_secret_examples(cloud_client: TestClient):
    from app.main import (
        create_app,  # noqa: F401  # openapi budowany niezaleznie od APP_ENV (zob. tools/export_openapi.py)
    )

    text = json.dumps(cloud_client.app.openapi())
    assert not re.search(r"daas_[0-9a-f]{8}_[A-Za-z0-9_-]{20,}", text)  # w przykladach nie ma pelnego klucza


def test_upload_path_is_never_taken_from_the_client(cloud_client: TestClient, cloud_settings: Settings, org):
    r = cloud_client.post(uploads_url(org.org_id), files={"file": ("a.csv", CSV_BYTES, "text/csv")},
                          data={"path": "/etc/passwd", "id": "../../x"}, headers=key_headers(org.api_key))
    assert r.status_code == 201
    (row,) = get_rows(cloud_settings, models.Upload)
    assert row.id == r.json()["id"] and len(row.id) == 32 and "/" not in row.path
    assert isinstance(Path(row.path).suffix, str)
